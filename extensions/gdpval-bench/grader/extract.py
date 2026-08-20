#!/usr/bin/env python3
"""Extract judge-readable text from a deliverable file.

Usage: extract.py <file> [max_chars]
Writes extracted text to stdout. Supports docx/xlsx/pptx/pdf/plain text.
Extraction is DATA handling only — the output is fed to judges wrapped in
untrusted-content fences.
"""
import sys
from pathlib import Path

MAX = int(sys.argv[2]) if len(sys.argv) > 2 else 40_000


def docx_text(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def xlsx_text(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=False, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"=== SHEET: {ws.title} ===")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 400:
                parts.append(f"... ({ws.max_row} rows total, truncated)")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells).rstrip())
    return "\n".join(parts)


def pptx_text(p):
    from pptx import Presentation
    prs = Presentation(p)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"=== SLIDE {i} ===")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(f"[notes] {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(parts)


def pdf_text(p):
    import pdfplumber
    parts = []
    with pdfplumber.open(p) as pdf:
        for i, page in enumerate(pdf.pages):
            if i >= 60:
                parts.append(f"... ({len(pdf.pages)} pages total, truncated)")
                break
            parts.append(f"=== PAGE {i + 1} ===")
            parts.append(page.extract_text() or "")
    return "\n".join(parts)


def main():
    p = Path(sys.argv[1])
    ext = p.suffix.lower()
    try:
        if ext in (".docx", ".docm"):
            text = docx_text(p)
        elif ext in (".xlsx", ".xlsm", ".xls"):
            text = xlsx_text(p)
        elif ext in (".pptx", ".pptm", ".ppt"):
            text = pptx_text(p)
        elif ext == ".pdf":
            text = pdf_text(p)
        else:
            text = p.read_text(errors="replace")
    except Exception as e:  # noqa: BLE001
        print(f"[extraction failed for {p.name}: {type(e).__name__}: {e}]")
        return
    if len(text) > MAX:
        text = text[:MAX] + f"\n... [truncated at {MAX} chars]"
    print(text)


main()
